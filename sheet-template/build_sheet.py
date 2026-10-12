"""Build the Topic Marketplace Google Sheet template (as .xlsx for import)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

NAVY, ORANGE, TEAL, GREEN = "00355F", "E37F1C", "5B9B98", "82A52C"
PLACEHOLDER = "FFF4C2"
OVER = "F4C7C3"

# Soft limits (warning, not reject). With per-card fit, last year's full-length topics
# (themes up to ~135 chars) across 13 rooms land at 21-23px at 1080p.
LIMITS = {"B": 110, "C": 140, "D": 110, "E": 140, "F": 60}

# HGI Riverwatch spaces as entered in the live Sheet 2026-10-12 ("- 1"/"- 2" = two conversations in one room).
ROOMS = ["Canal Room - 1", "Canal Room - 2", "Millpond Room - 1", "Millpond Room - 2",
         "Great Falls Room - 1", "Great Falls Room - 2", "Ballroom Rotunda", "Board Room",
         "Room 415", "Room 515", "Room 520", "Room 615", "Room 620"]
SPARE_ROWS = 7  # formatted and length-checked, left blank (blank Room = skipped by the display)

wb = Workbook()
ws = wb.active
ws.title = "Display"  # the display page reads this tab by name

headers = ["Room", "AM theme", "AM context", "PM theme", "PM context", "Note"]
fills = [NAVY, TEAL, TEAL, ORANGE, ORANGE, GREEN]
widths = [20, 42, 55, 42, 55, 32]
thin = Side(style="thin", color="C9D3DC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")

for i, (h, f, w) in enumerate(zip(headers, fills, widths), start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = PatternFill("solid", fgColor=f)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = border
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 24

last = len(ROOMS) + SPARE_ROWS + 1
for r, room in enumerate(ROOMS + [None] * SPARE_ROWS, start=2):
    a = ws.cell(row=r, column=1, value=room)
    a.font = Font(bold=True, color=NAVY)
    a.alignment = wrap_top
    a.border = border
    if room and "TBD" in room:
        a.fill = PatternFill("solid", fgColor=PLACEHOLDER)
    for col in range(2, 7):
        c = ws.cell(row=r, column=col)
        c.alignment = wrap_top
        c.border = border
    ws.row_dimensions[r].height = 48

ws.freeze_panes = "B2"

# Length checks: warning-style validation + red fill when over.
for col, n in LIMITS.items():
    rng = f"{col}2:{col}{last}"
    dv = DataValidation(
        type="custom",
        formula1=f"LEN({col}2)<={n}",
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="warning",
        errorTitle="Long for the projector",
        error=f"Over {n} characters. It will still show, but text shrinks for every room. Shorten if you can.",
        showInputMessage=True,
        promptTitle=f"Up to ~{n} characters",
        prompt="Context: one line per bullet (Ctrl+Enter or Alt+Enter for a new line)." if col in "CE"
        else ("Fills the whole card if both themes are blank (e.g. Reset Room: open all day)." if col == "F"
              else "Keep it short; it's read from the back of the room."),
    )
    dv.add(rng)
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f"LEN({col}2)>{n}"], fill=PatternFill("solid", fgColor=OVER))
    )

# ---------------- Instructions tab ----------------
ins = wb.create_sheet("How to use")
ins.column_dimensions["A"].width = 110
lines = [
    ("Topic Marketplace: room sheet", "title"),
    ("The display page reads the 'Display' tab every ~3 seconds and shows one card per row, in row order.", None),
    ("", None),
    ("Setup (once)", "h"),
    ("1. Upload this file to Google Drive, open it, then File > Save as Google Sheets. Work in the Google copy.", None),
    ("2. Share > General access > Anyone with the link > Viewer. (Topic text only; never put attendee names here.)", None),
    ("3. Copy the Sheet ID: the long string between /d/ and /edit in the URL. It goes into CONFIG.SHEET_ID on the display page.", None),
    ("4. Give edit access to the staff who will type on the day.", None),
    ("", None),
    ("Rules the display depends on", "h"),
    ("- Keep the tab named 'Display' and keep the header row and column order (Room, AM theme, AM context, PM theme, PM context, Note).", None),
    ("- Row order is display order. Reorder rooms by moving rows; the grid sizes itself to the room count (12 -> 4/4/4, 13 -> 5/4/4).", None),
    ("- A row with a blank Room is skipped, so you can leave spare rows at the bottom.", None),
    ("", None),
    ("Typing on the day", "h"),
    ("- Theme: the topic title. Context: optional supporting lines; each line in the cell becomes its own line on the card (Ctrl+Enter or Alt+Enter).", None),
    ("- Blank theme shows 'Topic coming soon', so unfilled rooms look intentional.", None),
    ("- Note: if both themes are blank, the Note fills the card (e.g. 'Reset Room: open all day', 'Not in use'). Otherwise it shows as a small line under the topics.", None),
    ("- Red cell = over the soft length limit. It still displays, but every card's text shrinks to fit, so shorten if you can.", None),
    (f"  Limits: themes {LIMITS['B']}, context {LIMITS['C']}, note {LIMITS['F']} characters.", None),
    ("- The projector laptop never opens this Sheet; type from a second device.", None),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = ins.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title":
        c.font = Font(bold=True, size=14, color=NAVY)
    elif kind == "h":
        c.font = Font(bold=True, size=11, color=ORANGE)

wb.active = 0
wb.save(OUT)
print("wrote", OUT)
